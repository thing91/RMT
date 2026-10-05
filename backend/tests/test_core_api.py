"""The core of the application end to end through the API: sign-in,
projects and documents, schema (item and link types), links, comments,
glossary, standards and reference documents, parameters, the traceability
matrix and saved views, the graph and exports."""

import io
import zipfile

import openpyxl


# ---- auth ---------------------------------------------------------------------


def test_login_provisions_user_and_token_works(world, client):
    r = client.post("/api/v1/auth/login", json={"username": "newbie", "password": "x"})
    assert r.status_code == 200, r.text
    token = r.json()["access_token"]
    me = client.get("/api/v1/auth/me", headers={"Authorization": f"Bearer {token}"})
    assert me.status_code == 200 and me.json()["username"] == "newbie"
    assert client.get("/api/v1/auth/me").status_code == 401
    assert client.get("/api/v1/auth/me", headers={"Authorization": "Bearer garbage"}).status_code == 401
    assert client.post("/api/v1/auth/login", json={"username": "newbie", "password": ""}).status_code == 401


# ---- projects, documents, groups ---------------------------------------------------


def test_projects_documents_and_groups(world):
    world.user("boss", "admin", org_admin=True)
    world.user("ed", "editor")
    world.user("outsider")
    c = world.client
    org = world.org.id

    r = c.post(f"/api/v1/organizations/{org}/projects", json={"name": "Rover", "visibility": "private", "key_prefix": "RVR"},
               headers=world.h("boss"))
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    assert c.post(f"/api/v1/organizations/{org}/projects", json={"name": "X", "key_prefix": "bad prefix"},
                  headers=world.h("boss")).status_code == 422

    # Private: invisible to someone without access, listed for the creator.
    names = [p["name"] for p in c.get(f"/api/v1/organizations/{org}/projects", headers=world.h("outsider")).json()]
    assert "Rover" not in names
    assert c.get(f"/api/v1/projects/{pid}", headers=world.h("outsider")).status_code == 404
    assert "Rover" in [p["name"] for p in c.get(f"/api/v1/organizations/{org}/projects", headers=world.h("boss")).json()]

    r = c.patch(f"/api/v1/projects/{pid}", json={"visibility": "public"}, headers=world.h("boss"))
    assert r.status_code == 200 and r.json()["visibility"] == "public"
    assert c.get(f"/api/v1/projects/{pid}").status_code == 200  # anonymous read of a public project

    g = c.post(f"/api/v1/projects/{pid}/document-groups", json={"name": "Specs"}, headers=world.h("boss")).json()
    d = c.post(f"/api/v1/projects/{pid}/documents", json={"name": "SRS", "group_id": g["id"]}, headers=world.h("boss"))
    assert d.status_code == 201, d.text
    docs = c.get(f"/api/v1/projects/{pid}/documents").json()
    assert [(x["name"], x["group_id"]) for x in docs] == [("SRS", g["id"])]

    assert c.patch(f"/api/v1/document-groups/{g['id']}", json={"name": "Specifications"},
                   headers=world.h("boss")).status_code == 200
    # Deleting a group ungroups its documents, it never deletes them.
    assert c.delete(f"/api/v1/document-groups/{g['id']}", headers=world.h("boss")).status_code == 204
    assert c.get(f"/api/v1/projects/{pid}/documents").json()[0]["group_id"] is None

    # Writes need a role on the project.
    assert c.post(f"/api/v1/projects/{pid}/documents", json={"name": "Nope"},
                  headers=world.h("outsider")).status_code in (403, 404)


def test_document_tree_numbers_sections_and_keys_requirements(world):
    world.user("ed", "editor")
    c = world.client
    s1 = world.create_item("ed", "General", type_=world.section)
    s11 = world.create_item("ed", "Mass", type_=world.section, parent_id=s1["id"])
    r1 = world.create_item("ed", "Mass ≤ 85 kg", parent_id=s11["id"])
    world.create_item("ed", "Power", type_=world.section)
    tree = c.get(f"/api/v1/documents/{world.doc.id}/items", headers=world.h("ed")).json()
    assert [n["number"] for n in tree] == ["1", "2"]
    assert tree[0]["children"][0]["number"] == "1.1"
    req = tree[0]["children"][0]["children"][0]
    assert req["id"] == r1["id"] and req["number"] is None and req["key"].startswith("SYS-")


# ---- schema: item types, attributes, link types -------------------------------------


def test_item_types_and_attribute_validation(world):
    world.user("boss", "admin", org_admin=True)
    world.user("ed", "editor")
    c = world.client
    org = world.org.id
    r = c.post(f"/api/v1/organizations/{org}/item-types",
               json={"key": "risk", "name": "Risk", "prefix": "RSK-", "behavior": "content"}, headers=world.h("boss"))
    assert r.status_code == 201, r.text
    risk = r.json()
    r = c.post(f"/api/v1/item-types/{risk['id']}/attributes", headers=world.h("boss"), json={
        "key": "likelihood", "label": "Likelihood", "data_type": "enum", "required": True,
        "enum_options": [{"value": "low", "label": "Low"}, {"value": "high", "label": "High"}]})
    assert r.status_code == 201, r.text
    c.post(f"/api/v1/item-types/{risk['id']}/attributes", headers=world.h("boss"),
           json={"key": "score", "label": "Score", "data_type": "number", "min": 0, "max": 10})
    # Editors can't change the schema.
    assert c.post(f"/api/v1/organizations/{org}/item-types", json={"key": "x", "name": "X", "prefix": "X-"},
                  headers=world.h("ed")).status_code == 403

    def make(**attrs):
        return c.post(f"/api/v1/documents/{world.doc.id}/items", headers=world.h("ed"),
                      json={"title": "Battery fire", "item_type_id": risk["id"], "attributes": attrs})

    assert make(likelihood="medium").status_code == 422  # not one of the options
    assert make(likelihood="low", score=11).status_code == 422  # above max
    ok = make(likelihood="high", score=7)
    assert ok.status_code == 201 and ok.json()["key"].startswith("RSK-")

    r = c.patch(f"/api/v1/item-types/{risk['id']}", json={"name": "Hazard"}, headers=world.h("boss"))
    assert r.status_code == 200 and r.json()["name"] == "Hazard"
    types = c.get(f"/api/v1/organizations/{org}/item-types").json()
    assert "Hazard" in {t["name"] for t in types}


def test_link_types_and_allowed_endpoints(world):
    world.user("boss", "admin", org_admin=True)
    world.user("ed", "editor")
    c = world.client
    r = c.post(f"/api/v1/organizations/{world.org.id}/link-types", headers=world.h("boss"), json={
        "key": "verified_by_test", "name": "Tested by", "inverse_name": "Tests",
        "allowed_source_types": [str(world.req.id)], "allowed_target_types": [str(world.va.id)]})
    assert r.status_code == 201, r.text
    lt = r.json()
    req = world.create_item("ed", "Req")
    va = world.create_item("ed", "Test", type_=world.va)
    bad = c.post("/api/v1/links", headers=world.h("ed"),
                 json={"source_item_id": va["id"], "target_item_id": req["id"], "link_type_id": lt["id"]})
    assert bad.status_code in (400, 422)
    good = c.post("/api/v1/links", headers=world.h("ed"),
                  json={"source_item_id": req["id"], "target_item_id": va["id"], "link_type_id": lt["id"]})
    assert good.status_code == 201, good.text
    dup = c.post("/api/v1/links", headers=world.h("ed"),
                 json={"source_item_id": req["id"], "target_item_id": va["id"], "link_type_id": lt["id"]})
    assert dup.status_code == 409
    links = c.get(f"/api/v1/items/{req['id']}/links", headers=world.h("ed")).json()
    assert [x["id"] for x in links] == [good.json()["id"]]
    assert len(c.get(f"/api/v1/projects/{world.project.id}/links", headers=world.h("ed")).json()) == 1


# ---- comments ----------------------------------------------------------------------


def test_comment_threads_reply_resolve_reopen_and_mentions(world):
    world.user("ed", "editor")
    world.user("rev", "reviewer")
    c = world.client
    item = world.create_item("ed", "Req")
    top = c.post(f"/api/v1/items/{item['id']}/comments", json={"body": "Units? @rev"}, headers=world.h("ed"))
    assert top.status_code == 201, top.text
    top = top.json()
    reply = c.post(f"/api/v1/items/{item['id']}/comments", json={"body": "kg", "parent_comment_id": top["id"]},
                   headers=world.h("rev"))
    assert reply.status_code == 201

    threads = c.get(f"/api/v1/items/{item['id']}/comments", headers=world.h("ed")).json()
    assert len(threads) == 1 and [r["body"] for r in threads[0]["replies"]] == ["kg"]
    # Mentioned → notified.
    notes = c.get("/api/v1/me/notifications", headers=world.h("rev")).json()["items"]
    assert any("mention" in (n.get("kind") or "") for n in notes)

    assert c.post(f"/api/v1/comments/{top['id']}/resolve", headers=world.h("ed")).status_code == 200
    doc_threads = c.get(f"/api/v1/documents/{world.doc.id}/comments", headers=world.h("ed")).json()
    assert doc_threads[0]["resolved"] is True
    assert c.post(f"/api/v1/comments/{top['id']}/reopen", headers=world.h("ed")).status_code == 200

    # Only the author edits their comment.
    assert c.patch(f"/api/v1/comments/{reply.json()['id']}", json={"body": "kilograms"},
                   headers=world.h("ed")).status_code == 403
    assert c.patch(f"/api/v1/comments/{reply.json()['id']}", json={"body": "kilograms"},
                   headers=world.h("rev")).status_code == 200
    assert c.delete(f"/api/v1/comments/{reply.json()['id']}", headers=world.h("rev")).status_code == 204


# ---- glossary ------------------------------------------------------------------------


def test_glossary_crud_report_and_xlsx_round_trip(world):
    world.user("boss", "admin", org_admin=True)
    world.user("ed", "editor")
    c = world.client
    org = world.org.id
    r = c.post(f"/api/v1/organizations/{org}/glossary", headers=world.h("boss"),
               json={"abbreviation": "EPS", "term": "Electrical Power System", "definition": "Generates and distributes power."})
    assert r.status_code == 201, r.text
    term = r.json()
    assert c.post(f"/api/v1/organizations/{org}/glossary", headers=world.h("ed"),
                  json={"abbreviation": "X", "term": "X", "definition": "x"}).status_code == 403
    assert c.get(f"/api/v1/organizations/{org}/glossary").json()[0]["abbreviation"] == "EPS"

    world.create_item("ed", "Power", description_md="The [[EPS]] shall use [[AOCS]] data.")
    report = c.get(f"/api/v1/organizations/{org}/glossary/consistency-report").json()
    assert "AOCS" in report["undefined_references"] and "EPS" not in report.get("unused_terms", [])

    xlsx = c.get(f"/api/v1/organizations/{org}/glossary/export.xlsx")
    assert xlsx.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
    assert "EPS" in [cell.value for cell in ws["A"]] + [cell.value for cell in ws["B"]]
    html = c.get(f"/api/v1/organizations/{org}/glossary/export.html")
    assert html.status_code == 200 and "Electrical Power System" in html.text

    # Import the same file back plus one new row → one added.
    wb = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    ws = wb.active
    header = [c_.value for c_ in ws[1]]
    row = {h: "" for h in header}
    for h in header:
        lower = str(h).lower()
        if "abbr" in lower:
            row[h] = "AOCS"
        elif lower.startswith("term"):
            row[h] = "Attitude and Orbit Control System"
        elif "defin" in lower:
            row[h] = "Points the spacecraft."
    ws.append([row[h] for h in header])
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post(f"/api/v1/organizations/{org}/glossary/import", headers=world.h("boss"),
               files={"file": ("g.xlsx", buf.getvalue(),
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    assert {t["abbreviation"] for t in c.get(f"/api/v1/organizations/{org}/glossary").json()} == {"EPS", "AOCS"}

    assert c.patch(f"/api/v1/glossary/{term['id']}", json={"definition": "Power."},
                   headers=world.h("boss")).status_code == 200
    assert c.delete(f"/api/v1/glossary/{term['id']}", headers=world.h("boss")).status_code == 204


# ---- standards & reference documents -----------------------------------------------------


def test_standards_locators_superseding_and_reference_documents(world, tmp_path, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "reference_files_dir", str(tmp_path))
    world.user("boss", "admin", org_admin=True)
    world.user("ed", "editor")
    c = world.client
    org = world.org.id
    old = c.post(f"/api/v1/organizations/{org}/standards", headers=world.h("boss"),
                 data={"code": "ECSS-E-10", "title": "System engineering", "version": "Rev. B", "source_type": "pdf"},
                 files={"file": ("e10.pdf", b"%PDF-1.4 fake", "application/pdf")})
    assert old.status_code == 201, old.text
    old = old.json()
    loc = c.post(f"/api/v1/reference-sources/{old['id']}/locators", json={"label": "5.4.2", "page_number": 12},
                 headers=world.h("boss"))
    assert loc.status_code == 201
    f = c.get(f"/api/v1/reference-sources/{old['id']}/file")
    assert f.status_code == 200 and f.content.startswith(b"%PDF")

    new = c.post(f"/api/v1/organizations/{org}/standards", headers=world.h("boss"),
                 data={"code": "ECSS-E-10C", "title": "System engineering", "version": "Rev. C",
                       "source_type": "link", "external_url": "https://example.com/e10c"}).json()
    world.create_item("ed", "Req", description_md="Per [[ref:ECSS-E-10#5.4.2]].")
    assert c.patch(f"/api/v1/reference-sources/{old['id']}", json={"superseded_by": new["id"]},
                   headers=world.h("boss")).status_code == 200
    report = c.get(f"/api/v1/organizations/{org}/standards/consistency-report").json()
    assert report["superseded_citations"] and "ECSS-E-10C" in report["unused_sources"]

    ref = c.post(f"/api/v1/projects/{world.project.id}/reference-documents", headers=world.h("ed"),
                 data={"code": "TN-01", "title": "Power budget", "source_type": "link",
                       "external_url": "https://example.com/tn01"})
    assert ref.status_code == 201, ref.text
    assert [r["code"] for r in c.get(f"/api/v1/projects/{world.project.id}/reference-documents",
                                    headers=world.h("ed")).json()] == ["TN-01"]
    assert c.delete(f"/api/v1/reference-locators/{loc.json()['id']}", headers=world.h("boss")).status_code == 204


# ---- parameters -------------------------------------------------------------------------


def test_parameters_constant_derived_cycle_and_usage(world):
    world.user("ed", "editor")
    c = world.client
    pid = world.project.id

    def create(**body):
        return c.post(f"/api/v1/projects/{pid}/parameters", json=body, headers=world.h("ed"))

    assert create(key="MASS", name="Mass", value=80, unit="kg", data_type="constant").status_code == 201
    assert create(key="G", name="Gravity", value=9.81, unit="m/s2", data_type="constant").status_code == 201
    w = create(key="WEIGHT", name="Weight", unit="N", data_type="derived", formula="MASS * G")
    assert w.status_code == 201, w.text
    params = {p["key"]: p for p in c.get(f"/api/v1/projects/{pid}/parameters", headers=world.h("ed")).json()}
    assert abs(params["WEIGHT"]["computed_value"] - 784.8) < 1e-6

    # No arbitrary code in formulas, and no cycles.
    assert create(key="EVIL", name="x", data_type="derived", formula="__import__('os')").status_code == 422
    bad = c.patch(f"/api/v1/parameters/{params['MASS']['id']}", json={"data_type": "derived", "formula": "WEIGHT / G"},
                  headers=world.h("ed"))
    assert bad.status_code == 422

    # The rejected change was not saved.
    params = {p["key"]: p for p in c.get(f"/api/v1/projects/{pid}/parameters", headers=world.h("ed")).json()}
    assert params["MASS"]["data_type"] == "constant" and params["MASS"]["computed_value"] == 80
    assert "EVIL" not in params
    assert create(key="SELF", name="x", data_type="derived", formula="SELF + 1").status_code == 422
    assert create(key="TYPO", name="x", data_type="derived", formula="MAS * 2").status_code == 422

    # A derived parameter built on another derived one (regression: it used
    # to resolve to nothing).
    r = create(key="WEIGHT_MARGIN", name="Weight with margin", data_type="derived", formula="WEIGHT * 1.1")
    assert r.status_code == 201, r.text
    assert abs(r.json()["computed_value"] - 863.28) < 1e-6

    item = world.create_item("ed", "Weight limit", description_md="Weight below {{param:WEIGHT}}.")
    usages = c.get(f"/api/v1/parameters/{params['WEIGHT']['id']}/usages", headers=world.h("ed")).json()
    assert [u["item_id"] for u in usages] == [item["id"]]


# ---- traceability matrix, saved views, graph -----------------------------------------------


def test_matrix_saved_views_table_views_and_graph(world):
    world.user("ed", "editor")
    world.user("other", "editor")
    c = world.client
    pid = world.project.id
    r1 = world.create_item("ed", "R1")
    world.create_item("ed", "R2")
    v1 = world.create_item("ed", "V1", type_=world.va)
    c.post("/api/v1/links", headers=world.h("ed"),
           json={"source_item_id": v1["id"], "target_item_id": r1["id"], "link_type_id": str(world.verifies.id)})

    m = c.get(f"/api/v1/projects/{pid}/traceability-matrix", headers=world.h("ed"), params={
        "row_type_id": str(world.req.id), "col_type_id": str(world.va.id), "link_type_id": str(world.verifies.id)})
    assert m.status_code == 200, m.text
    body = m.json()
    assert len(body["rows"]) == 2 and len(body["cols"]) == 1
    assert len(body.get("uncovered_row_ids", body.get("uncovered", []))) == 1

    personal = c.post(f"/api/v1/projects/{pid}/matrix-views", headers=world.h("ed"), json={
        "name": "Coverage", "scope": "personal", "row_type_id": str(world.req.id),
        "col_type_id": str(world.va.id), "link_type_id": str(world.verifies.id)})
    assert personal.status_code == 201, personal.text
    assert c.get(f"/api/v1/projects/{pid}/matrix-views", headers=world.h("other")).json() == []
    assert len(c.get(f"/api/v1/projects/{pid}/matrix-views", headers=world.h("ed")).json()) == 1

    tv = c.post(f"/api/v1/projects/{pid}/table-views", headers=world.h("ed"), json={
        "name": "Owners", "item_type_id": str(world.req.id), "columns": [{"attribute_key": "owner"}]})
    assert tv.status_code == 201, tv.text
    upd = c.patch(f"/api/v1/table-views/{tv.json()['id']}", json={"name": "By owner"}, headers=world.h("ed"))
    assert upd.status_code == 200 and upd.json()["name"] == "By owner"
    assert c.delete(f"/api/v1/table-views/{tv.json()['id']}", headers=world.h("ed")).status_code == 204

    g = c.get(f"/api/v1/items/{r1['id']}/graph", headers=world.h("ed"))
    assert g.status_code == 200, g.text
    ids = {n["id"] for n in g.json()["nodes"]}
    assert {r1["id"], v1["id"]} <= ids and len(g.json()["edges"]) == 1


# ---- exports ------------------------------------------------------------------------------


def test_document_exports(world):
    world.user("ed", "editor")
    c = world.client
    sec = world.create_item("ed", "General", type_=world.section)
    world.create_item("ed", "Mass ≤ 85 kg", parent_id=sec["id"], description_md="Wet mass incl. margin.")

    x = c.get(f"/api/v1/documents/{world.doc.id}/export.xlsx", headers=world.h("ed"))
    assert x.status_code == 200
    values = [cell.value for row in openpyxl.load_workbook(io.BytesIO(x.content)).active.iter_rows() for cell in row]
    assert "Mass ≤ 85 kg" in values

    h = c.get(f"/api/v1/documents/{world.doc.id}/export.html", headers=world.h("ed"))
    assert h.status_code == 200 and "Mass ≤ 85 kg" in h.text and "General" in h.text

    q = c.get(f"/api/v1/documents/{world.doc.id}/export.reqif", headers=world.h("ed"))
    assert q.status_code == 200
    content = q.content
    if content[:2] == b"PK":
        content = zipfile.ZipFile(io.BytesIO(content)).read(zipfile.ZipFile(io.BytesIO(content)).namelist()[0])
    assert b"REQ-IF" in content and "Mass ≤ 85 kg".encode() in content

    s = c.get(f"/api/v1/projects/{world.project.id}/suggestions/export.xlsx", headers=world.h("ed"))
    assert s.status_code == 200


def test_health_and_git_log(client, world):
    world.user("ed", "editor")
    assert client.get("/api/v1/health").json() == {"status": "ok"}
    r = client.get(f"/api/v1/projects/{world.project.id}/git-log", headers=world.h("ed"))
    assert r.status_code == 200 and r.json() == []  # git sync is disabled in tests
