"""Generic requirements library, releases, applicability matrix, updates,
exports, parametric requirements and rule-based suggestions."""

import pytest


@pytest.fixture
def lib_world(world):
    from app.models.document import Document
    from app.models.project import Project
    from app.models.project_access import ProjectAccess

    db = world.db
    library = Project(org_id=world.org.id, name="GEN-SAT", visibility="private", kind="library",
                      mission_attribute_defs=[{"key": "propulsion", "label": "Propulsion", "type": "boolean"}])
    db.add(library)
    db.flush()
    eps = Document(project_id=library.id, name="EPS", order=0)
    aocs = Document(project_id=library.id, name="AOCS", order=1)
    db.add_all([eps, aocs])
    db.commit()
    world.user("libadmin", None)
    world.user("pm", "admin")  # project admin of the consuming project
    world.user("eng", "editor")
    world.user("pa", "editor")
    for username in ("libadmin",):
        db.add(ProjectAccess(project_id=library.id, subject_type="user", subject_id=world.users[username].id, role="admin"))
    db.commit()
    world.library = library
    world.eps = eps
    world.aocs = aocs
    return world


def _lib_item(w, doc, title, description=None):
    r = w.client.post(
        f"/api/v1/documents/{doc.id}/items",
        json={"title": title, "item_type_id": str(w.req.id), "description_md": description},
        headers=w.h("libadmin"),
    )
    assert r.status_code == 201, r.text
    return r.json()


def _release(w, name):
    r = w.client.post(f"/api/v1/projects/{w.library.id}/releases", json={"name": name}, headers=w.h("libadmin"))
    assert r.status_code == 201, r.text
    return r.json()


def _subscribe(w, release):
    r = w.client.post(
        f"/api/v1/projects/{w.project.id}/libraries",
        json={"library_id": str(w.library.id), "release_id": release["id"], "target_document_id": str(w.doc.id)},
        headers=w.h("pm"),
    )
    assert r.status_code == 201, r.text
    return r.json()


def _records(w, user="eng", **params):
    r = w.client.get(f"/api/v1/projects/{w.project.id}/applicability", params=params, headers=w.h(user))
    assert r.status_code == 200, r.text
    return r.json()


def _run(w, user, record, key, **body):
    return w.client.post(f"/api/v1/workflow/applicability/{record['id']}/transitions/{key}", json=body,
                         headers=w.h(user))


def test_import_creates_references_not_copies(lib_world):
    w = lib_world
    _lib_item(w, w.eps, "Battery sized for eclipse", "Battery shall supply power for {{param:eclipse_duration}}.")
    _lib_item(w, w.aocs, "Pointing accuracy")
    rel = _release(w, "GEN-SAT v1.0")
    assert rel["item_count"] == 2
    sub = _subscribe(w, rel)
    assert sub["release"]["name"] == "GEN-SAT v1.0"

    recs = _records(w)
    assert {r["subsystem"] for r in recs} == {"EPS", "AOCS"}
    assert all(r["status"] == "to_be_assessed" for r in recs)
    # No item was copied into the project.
    tree = w.client.get(f"/api/v1/documents/{w.doc.id}/items", headers=w.h("eng")).json()
    assert tree == []
    summary = w.client.get(f"/api/v1/projects/{w.project.id}/applicability/summary", headers=w.h("eng")).json()
    assert summary["complete"] is False and summary["totals"]["to_be_assessed"] == 2


def test_decisions_and_completeness(lib_world):
    w = lib_world
    _lib_item(w, w.eps, "Battery", "Battery shall last.")
    _lib_item(w, w.aocs, "Thruster control", "Thrusters shall be controlled.")
    _lib_item(w, w.aocs, "Pointing", "Point accurately.")
    _subscribe(w, _release(w, "v1"))
    recs = {r["title"]: r for r in _records(w)}

    # 1:1
    assert _run(w, "eng", recs["Battery"], "mark_applicable").status_code == 200
    # Not applicable needs a justification…
    assert _run(w, "eng", recs["Thruster control"], "propose_not_applicable").status_code == 422
    assert _run(w, "eng", recs["Thruster control"], "propose_not_applicable",
                comment="No propulsion on this mission").status_code == 200
    # …and approval by someone with applicability.approve (not the proposer).
    r = _run(w, "pa", recs["Thruster control"], "approve_not_applicable")
    assert r.status_code == 403
    assert _run(w, "pm", recs["Thruster control"], "approve_not_applicable").status_code == 200

    # With modification → a project copy appears, with a diff vs. the library.
    r = _run(w, "eng", recs["Pointing"], "mark_modified", comment="Tighter accuracy needed",
             payload={"description_md": "Point accurately within 0.1 deg."})
    assert r.status_code == 200, r.text
    rec = w.client.get(f"/api/v1/applicability/{recs['Pointing']['id']}", headers=w.h("eng")).json()
    assert rec["modified_item_key"] and rec["rationale"] == "Tighter accuracy needed"
    assert rec["modification_diff"][0]["field"] == "description_md"
    copy = w.client.get(f"/api/v1/items/{rec['modified_item_id']}", headers=w.h("eng")).json()
    assert copy["source_item_id"] == recs["Pointing"]["source_item_id"]
    assert copy["attributes"]["status"] == "draft"  # goes through the normal requirement workflow

    summary = w.client.get(f"/api/v1/projects/{w.project.id}/applicability/summary", headers=w.h("eng")).json()
    assert summary["complete"] is True
    assert summary["by_subsystem"]["AOCS"]["not_applicable"] == 1

    audit = w.client.get(f"/api/v1/projects/{w.project.id}/audit?action=workflow", headers=w.h("pm")).json()
    assert any(e["subject_type"] == "applicability" and e["payload"].get("comment") == "No propulsion on this mission"
               for e in audit)


def test_library_update_flow(lib_world):
    w = lib_world
    a = _lib_item(w, w.eps, "A", "Alpha")
    _lib_item(w, w.eps, "B", "Bravo")
    c = _lib_item(w, w.eps, "C", "Charlie")
    _subscribe(w, _release(w, "v1"))
    recs = {r["title"]: r for r in _records(w)}
    _run(w, "eng", recs["A"], "mark_applicable")
    _run(w, "eng", recs["C"], "mark_applicable")

    # Library evolves: A changes, C is retired, D is new.
    w.patch("libadmin", a, description_md="Alpha, revised")
    w.transition("libadmin", c, "retire", comment="Superseded")
    _lib_item(w, w.eps, "D", "Delta")
    rel2 = _release(w, "v2")

    subs = w.client.get(f"/api/v1/projects/{w.project.id}/libraries", headers=w.h("pm")).json()
    assert subs[0]["update_available"] is True
    preview = w.client.get(
        f"/api/v1/projects/{w.project.id}/libraries/{w.library.id}/update-preview",
        params={"release_id": rel2["id"]}, headers=w.h("pm"),
    ).json()
    assert [x["title"] for x in preview["added"]] == ["D"]
    assert [x["title"] for x in preview["changed"]] == ["A"]
    assert [x["title"] for x in preview["removed"]] == ["C"]

    # Staying needs a justification.
    r = w.client.post(f"/api/v1/projects/{w.project.id}/libraries/{w.library.id}/decline",
                      json={"release_id": rel2["id"], "justification": " "}, headers=w.h("pm"))
    assert r.status_code == 422

    r = w.client.post(f"/api/v1/projects/{w.project.id}/libraries/{w.library.id}/adopt",
                      json={"release_id": rel2["id"]}, headers=w.h("pm"))
    assert r.status_code == 200, r.text
    recs = {r["title"]: r for r in _records(w)}
    assert recs["A"]["suspect"] and "Changed in v2" in recs["A"]["suspect_reason"]
    assert recs["A"]["description_md"] == "Alpha, revised"
    assert recs["C"]["removed_from_library"] and recs["C"]["suspect"]
    assert recs["D"]["status"] == "to_be_assessed"
    assert recs["B"]["suspect"] is False

    r = w.client.post(f"/api/v1/applicability/{recs['A']['id']}/clear-suspect", headers=w.h("eng"))
    assert r.status_code == 200 and r.json()["suspect"] is False


def test_parametric_requirements(lib_world):
    from app.models.parameter import Parameter

    w = lib_world
    _lib_item(w, w.eps, "Eclipse", "Battery shall supply power for {{param:eclipse_duration}}.")
    _subscribe(w, _release(w, "v1"))
    rec = _records(w)[0]
    assert rec["missing_parameters"] == ["eclipse_duration"]
    w.db.add(Parameter(project_id=w.project.id, key="eclipse_duration", name="Eclipse", value=35, unit="min",
                       data_type="constant", source="assumed"))
    w.db.commit()
    rec = _records(w)[0]
    assert rec["resolved_description"] == "Battery shall supply power for 35 min."
    # Setting a parameter isn't a modification: 1:1 still possible.
    assert _run(w, "eng", rec, "mark_applicable").status_code == 200


def test_rule_based_suggestion(lib_world):
    w = lib_world
    thr = _lib_item(w, w.aocs, "Thruster control")
    r = w.client.put(
        f"/api/v1/items/{thr['id']}/library-rules",
        json={"rules": [{"attribute": "propulsion", "op": "falsy", "suggest": "not_applicable",
                         "reason": "Mission has no propulsion"}]},
        headers=w.h("libadmin"),
    )
    assert r.status_code == 200, r.text
    _subscribe(w, _release(w, "v1"))
    assert _records(w)[0]["suggested_status"] is None  # attribute not set yet

    r = w.client.put(f"/api/v1/projects/{w.project.id}/mission-attributes", json={"values": {"propulsion": False}},
                     headers=w.h("pm"))
    assert r.json()["suggestions"] == 1
    rec = _records(w)[0]
    assert rec["suggested_status"] == "not_applicable" and rec["status"] == "to_be_assessed"


def test_exports(lib_world):
    from io import BytesIO
    from xml.etree import ElementTree as ET

    from openpyxl import load_workbook

    w = lib_world
    _lib_item(w, w.eps, "Battery", "Battery shall last.")
    _subscribe(w, _release(w, "v1"))
    r = w.client.get(f"/api/v1/projects/{w.project.id}/applicability/export.xlsx", headers=w.h("eng"))
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb["Applicability"]["F2"].value == "To be assessed"
    assert wb["Summary"]["A1"].value == "Subsystem"

    r = w.client.get(f"/api/v1/projects/{w.project.id}/applicability/export.reqif", headers=w.h("eng"))
    assert r.status_code == 200
    root = ET.fromstring(r.content)
    ns = {"r": "http://www.omg.org/spec/ReqIF/20110401/reqif.xsd"}
    assert len(root.findall(".//r:SPEC-OBJECT", ns)) == 1
    assert root.find(".//r:ENUM-VALUE-REF", ns).text == "EV-to_be_assessed"


def test_review_can_require_complete_matrix(lib_world):
    w = lib_world
    _lib_item(w, w.eps, "Battery")
    _subscribe(w, _release(w, "v1"))
    item = w.create_item("eng", "Project req")
    r = w.client.post(
        f"/api/v1/projects/{w.project.id}/reviews",
        json={"title": "SRR", "item_ids": [item["id"]], "participants": [], "require_applicability_complete": True},
        headers=w.h("pm"),
    )
    review = r.json()
    assert "applicability matrix" in review["close_blockers"][0]


def test_baseline_freezes_applicability(lib_world):
    w = lib_world
    _lib_item(w, w.eps, "Battery")
    _subscribe(w, _release(w, "v1"))
    r = w.client.post(f"/api/v1/projects/{w.project.id}/baselines", json={"name": "SRR"}, headers=w.h("pm"))
    assert r.status_code == 201, r.text
    detail = w.client.get(f"/api/v1/baselines/{r.json()['id']}", headers=w.h("pm")).json()
    assert detail["snapshot"]["applicability"][0]["status"] == "to_be_assessed"
