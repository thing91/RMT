"""Change suggestions API (komentarze i sugestie zmian)."""

import io
import uuid
from datetime import datetime
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, get_db, require_user
from app.models.comment import ItemComment
from app.models.item import Item
from app.models.project import Project
from app.models.suggestion import SUGGESTION_KINDS, SUGGESTION_SEVERITIES, Suggestion
from app.models.user import User
from app.models.workflow_def import Review
from app.services import suggestions as svc
from app.services import workflow_engine as wf
from app.services.permissions import can_read
from app.services.versioning import _text_ops
from app.services.workflow_templates import SYSTEM_WORKFLOWS

router = APIRouter(tags=["suggestions"])


class SuggestionCreate(BaseModel):
    target_field: str
    proposed_value: Any = None
    rationale: str
    kind: str = "technical"
    severity: str = "minor"
    review_id: uuid.UUID | None = None


class SuggestionUpdate(BaseModel):
    proposed_value: Any = None
    rationale: str | None = None
    kind: str | None = None
    severity: str | None = None


class ConvertComment(BaseModel):
    target_field: str
    proposed_value: Any = None
    rationale: str | None = None  # defaults to the comment body
    kind: str = "technical"
    severity: str = "minor"
    review_id: uuid.UUID | None = None


class SuggestionRead(BaseModel):
    id: uuid.UUID
    key: str
    project_id: uuid.UUID
    item_id: uuid.UUID
    item_key: str | None
    item_title: str
    document_id: uuid.UUID
    item_version: int
    item_current_version: int
    target_field: str
    target_label: str
    base_value: Any
    proposed_value: Any
    final_value: Any
    current_value: Any
    rationale: str
    kind: str
    severity: str
    status: str
    status_label: str
    review_id: uuid.UUID | None
    review_key: str | None
    source_comment_id: uuid.UUID | None
    author_id: uuid.UUID
    author: str
    resolution_text: str | None
    resolved_by: str | None
    resolved_at: datetime | None
    apply_pending: bool
    applied_version: int | None
    created_at: datetime
    updated_at: datetime
    # Word-level diff of the proposal against the value it was based on,
    # and (when Outdated) against the item's current value.
    diff: list[dict] | None = None
    diff_vs_current: list[dict] | None = None


def _target_label(s: Suggestion) -> str:
    f = s.target_field
    if f == "title":
        return "Title"
    if f == "description_md":
        return "Description"
    if f == "links.add":
        return "Add link"
    if f == "links.remove":
        return "Remove link"
    key = f.split(".", 1)[1]
    attr = next((a for a in s.item.item_type.attributes if a.key == key), None)
    return attr.label if attr else key


def _as_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return str(v)


def _read(db: Session, s: Suggestion, with_diff: bool = True) -> SuggestionRead:
    current = svc.current_value(db, s.item, s.target_field)
    textual = s.target_field not in svc.LINK_FIELDS
    return SuggestionRead(
        id=s.id,
        key=s.key,
        project_id=s.project_id,
        item_id=s.item_id,
        item_key=s.item.key,
        item_title=s.item.title,
        document_id=s.item.document_id,
        item_version=s.item_version,
        item_current_version=s.item.version,
        target_field=s.target_field,
        target_label=_target_label(s),
        base_value=s.base_value,
        proposed_value=s.proposed_value,
        final_value=s.final_value,
        current_value=current if textual else None,
        rationale=s.rationale,
        kind=s.kind,
        severity=s.severity,
        status=s.status,
        status_label=wf.state_label(SYSTEM_WORKFLOWS["suggestion"], s.status),
        review_id=s.review_id,
        review_key=s.review.key if s.review else None,
        source_comment_id=s.source_comment_id,
        author_id=s.author_id,
        author=s.author.display_name,
        resolution_text=s.resolution_text,
        resolved_by=s.resolved_by.display_name if s.resolved_by else None,
        resolved_at=s.resolved_at,
        apply_pending=s.apply_pending,
        applied_version=s.applied_version,
        created_at=s.created_at,
        updated_at=s.updated_at,
        diff=_text_ops(_as_text(s.base_value), _as_text(s.proposed_value)) if (with_diff and textual) else None,
        diff_vs_current=(
            _text_ops(_as_text(current), _as_text(s.proposed_value))
            if (with_diff and textual and s.status == "outdated")
            else None
        ),
    )


def _item(db: Session, item_id: uuid.UUID, user: User | None) -> Item:
    item = db.get(Item, item_id)
    if not item or not can_read(db, item.document.project, user.id if user else None):
        raise HTTPException(404, "Item not found")
    return item


def _suggestion(db: Session, suggestion_id: uuid.UUID, user: User | None) -> Suggestion:
    s = db.get(Suggestion, suggestion_id)
    if not s:
        raise HTTPException(404, "Suggestion not found")
    _item(db, s.item_id, user)
    return s


def _check_classification(kind: str | None, severity: str | None) -> None:
    if kind is not None and kind not in SUGGESTION_KINDS:
        raise HTTPException(422, f"kind must be one of {SUGGESTION_KINDS}")
    if severity is not None and severity not in SUGGESTION_SEVERITIES:
        raise HTTPException(422, f"severity must be one of {SUGGESTION_SEVERITIES}")


def _check_review(db: Session, review_id: uuid.UUID | None, item: Item) -> None:
    if review_id is None:
        return
    review = db.get(Review, review_id)
    if not review or review.project_id != item.document.project_id:
        raise HTTPException(422, "Review not found in this project.")
    if review.status != "open":
        raise HTTPException(409, "That review is closed.")


@router.get("/items/{item_id}/suggestions", response_model=list[SuggestionRead])
def list_item_suggestions(
    item_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)
):
    _item(db, item_id, user)
    rows = db.query(Suggestion).filter(Suggestion.item_id == item_id).order_by(Suggestion.created_at.desc()).all()
    return [_read(db, s) for s in rows]


@router.post("/items/{item_id}/suggestions", response_model=SuggestionRead, status_code=201)
def create_suggestion(
    item_id: uuid.UUID, payload: SuggestionCreate, user: User = Depends(require_user), db: Session = Depends(get_db)
):
    """Anyone who can read the item may suggest — no edit permission and no
    edit lock needed, and any number of people can do it in parallel."""
    item = _item(db, item_id, user)
    _check_classification(payload.kind, payload.severity)
    _check_review(db, payload.review_id, item)
    try:
        s = svc.create(
            db, item, user, target_field=payload.target_field, proposed_value=payload.proposed_value,
            rationale=payload.rationale, kind=payload.kind, severity=payload.severity, review_id=payload.review_id,
        )
    except svc.SuggestionError as e:
        raise HTTPException(422, str(e))
    db.commit()
    db.refresh(s)
    return _read(db, s)


@router.post("/comments/{comment_id}/convert-to-suggestion", response_model=SuggestionRead, status_code=201)
def convert_comment(
    comment_id: uuid.UUID, payload: ConvertComment, user: User = Depends(require_user), db: Session = Depends(get_db)
):
    comment = db.get(ItemComment, comment_id)
    if not comment:
        raise HTTPException(404, "Comment not found")
    item = _item(db, comment.item_id, user)
    _check_classification(payload.kind, payload.severity)
    _check_review(db, payload.review_id, item)
    try:
        s = svc.create(
            db, item, user, target_field=payload.target_field, proposed_value=payload.proposed_value,
            rationale=payload.rationale or comment.body, kind=payload.kind, severity=payload.severity,
            review_id=payload.review_id, source_comment_id=comment.id,
        )
    except svc.SuggestionError as e:
        raise HTTPException(422, str(e))
    db.commit()
    db.refresh(s)
    return _read(db, s)


@router.get("/suggestions/{suggestion_id}", response_model=SuggestionRead)
def get_suggestion(
    suggestion_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)
):
    return _read(db, _suggestion(db, suggestion_id, user))


@router.patch("/suggestions/{suggestion_id}", response_model=SuggestionRead)
def update_suggestion(
    suggestion_id: uuid.UUID,
    payload: SuggestionUpdate,
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    """The author may revise an open or outdated suggestion. Revising an
    outdated one rebases it on the item's current version and reopens it."""
    s = _suggestion(db, suggestion_id, user)
    if s.author_id != user.id:
        raise HTTPException(403, "Only the author can revise a suggestion.")
    if s.status not in ("open", "outdated"):
        raise HTTPException(409, "Only open or outdated suggestions can be revised.")
    data = payload.model_dump(exclude_unset=True)
    _check_classification(data.get("kind"), data.get("severity"))
    if "proposed_value" in data:
        try:
            svc.validate_target(s.item, s.target_field, data["proposed_value"])
        except svc.SuggestionError as e:
            raise HTTPException(422, str(e))
    if "rationale" in data and not (data["rationale"] or "").strip():
        raise HTTPException(422, "A justification is required.")
    for k, v in data.items():
        setattr(s, k, v)
    if s.status == "outdated":
        svc.rebase(db, s)
        wf.execute(db, svc.SuggestionSubject(db, s), "refresh", user,
                   comment=f"Rebased on version {s.item.version}.")
    db.commit()
    db.refresh(s)
    return _read(db, s)


@router.get("/projects/{project_id}/suggestions", response_model=list[SuggestionRead])
def list_project_suggestions(
    project_id: uuid.UUID,
    review_id: uuid.UUID | None = None,
    item_id: uuid.UUID | None = None,
    author_id: uuid.UUID | None = None,
    status: str | None = None,
    kind: str | None = None,
    severity: str | None = None,
    user: User | None = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project = db.get(Project, project_id)
    if not project or not can_read(db, project, user.id if user else None):
        raise HTTPException(404, "Project not found")
    q = _filtered(db, project_id, review_id, item_id, author_id, status, kind, severity)
    return [_read(db, s, with_diff=False) for s in q.all()]


def _filtered(db, project_id, review_id, item_id, author_id, status, kind, severity):
    q = db.query(Suggestion).filter(Suggestion.project_id == project_id)
    if review_id:
        q = q.filter(Suggestion.review_id == review_id)
    if item_id:
        q = q.filter(Suggestion.item_id == item_id)
    if author_id:
        q = q.filter(Suggestion.author_id == author_id)
    if status:
        q = q.filter(Suggestion.status.in_(status.split(",")))
    if kind:
        q = q.filter(Suggestion.kind == kind)
    if severity:
        q = q.filter(Suggestion.severity == severity)
    return q.order_by(Suggestion.key)


_DECISION = {
    "accepted": "Accepted",
    "accepted_modified": "Accepted with modification",
    "rejected": "Rejected",
    "withdrawn": "Withdrawn",
    "outdated": "Outdated",
    "open": "Open",
}


@router.get("/projects/{project_id}/suggestions/export.xlsx")
def export_suggestions(
    project_id: uuid.UUID,
    review_id: uuid.UUID | None = None,
    status: str | None = None,
    kind: str | None = None,
    severity: str | None = None,
    user: User | None = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """RID table for the customer: ID, requirement, proposed change,
    justification, classification, decision, response."""
    project = db.get(Project, project_id)
    if not project or not can_read(db, project, user.id if user else None):
        raise HTTPException(404, "Project not found")
    rows = _filtered(db, project_id, review_id, None, None, status, kind, severity).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "RIDs"
    header = [
        "ID", "Requirement", "Requirement title", "Field", "Current text", "Proposed change",
        "Justification", "Type", "Severity", "Author", "Review", "Decision", "Response", "Decided by", "Decided at",
    ]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for s in rows:
        ws.append(
            [
                s.key,
                s.item.key or "",
                s.item.title,
                _target_label(s),
                _as_text(s.base_value),
                _as_text(s.final_value if s.final_value is not None else s.proposed_value),
                s.rationale,
                "Editorial" if s.kind == "editorial" else "Technical",
                s.severity.title(),
                s.author.display_name,
                s.review.key if s.review else "",
                _DECISION.get(s.status, s.status),
                s.resolution_text or "",
                s.resolved_by.display_name if s.resolved_by else "",
                s.resolved_at.strftime("%Y-%m-%d") if s.resolved_at else "",
            ]
        )
    widths = [11, 12, 30, 14, 40, 40, 40, 11, 9, 16, 10, 22, 40, 16, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"RIDs-{project.name}.xlsx".replace(" ", "_")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/suggestions/{suggestion_id}/comments")
def suggestion_comments(
    suggestion_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)
):
    from app.api.v1.comments import _threads

    s = _suggestion(db, suggestion_id, user)
    return _threads(db.query(ItemComment).filter(ItemComment.suggestion_id == s.id).all())
