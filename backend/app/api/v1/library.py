"""Generic requirements library + applicability matrix API."""

import io
import uuid
from datetime import datetime
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, get_db, require_user
from app.models.item import Item
from app.models.library import ApplicabilityRecord, LibraryRelease, LibraryRule, ProjectLibrary
from app.models.project import Project
from app.models.user import User
from app.services import audit
from app.services import library as lib
from app.services import workflow_engine as wf
from app.services.permissions import can_read, has_org_permission, has_permission
from app.services.workflow_templates import SYSTEM_WORKFLOWS

router = APIRouter(tags=["library"])


# ---- schemas -------------------------------------------------------------------


class ReleaseCreate(BaseModel):
    name: str
    notes: str | None = None


class ReleaseRead(BaseModel):
    id: uuid.UUID
    library_id: uuid.UUID
    name: str
    notes: str | None
    item_count: int
    released_by: str | None
    released_at: datetime
    previous_release_id: uuid.UUID | None


class LibraryRead(BaseModel):
    id: uuid.UUID
    name: str
    latest_release: ReleaseRead | None
    requirement_count: int
    mission_attribute_defs: list


class SubscribeIn(BaseModel):
    library_id: uuid.UUID
    release_id: uuid.UUID
    target_document_id: uuid.UUID | None = None


class AdoptIn(BaseModel):
    release_id: uuid.UUID


class DeclineIn(BaseModel):
    release_id: uuid.UUID
    justification: str


class SubscriptionPatch(BaseModel):
    target_document_id: uuid.UUID | None = None


class SubscriptionRead(BaseModel):
    library_id: uuid.UUID
    library_name: str
    release: ReleaseRead
    latest_release: ReleaseRead | None
    update_available: bool
    declined_release: ReleaseRead | None
    declined_justification: str | None
    target_document_id: uuid.UUID | None


class RecordRead(BaseModel):
    id: uuid.UUID
    library_id: uuid.UUID
    release_name: str
    source_item_id: uuid.UUID
    source_version: int
    key: str | None
    title: str
    description_md: str | None
    resolved_title: str
    resolved_description: str
    missing_parameters: list[str]
    subsystem: str
    status: str
    status_label: str
    rationale: str | None
    modified_item_id: uuid.UUID | None
    modified_item_key: str | None
    modified_document_id: uuid.UUID | None
    decided_by: str | None
    decided_at: datetime | None
    approved_by: str | None
    approved_at: datetime | None
    suspect: bool
    suspect_reason: str | None
    removed_from_library: bool
    suggested_status: str | None
    suggested_reason: str | None


class RecordDetail(RecordRead):
    modification_diff: list[dict]
    attributes: dict


class MissionAttributesIn(BaseModel):
    values: dict[str, Any]


class MissionDefsIn(BaseModel):
    defs: list[dict]


class RulesIn(BaseModel):
    rules: list[dict]


# ---- helpers ---------------------------------------------------------------------


def _project(db: Session, project_id: uuid.UUID, user: User | None) -> Project:
    project = db.get(Project, project_id)
    if not project or not can_read(db, project, user.id if user else None):
        raise HTTPException(404, "Project not found")
    return project


def _require(db: Session, project_id: uuid.UUID, user: User | None, perm: str) -> User:
    if user is None or not has_permission(db, project_id, user, perm):
        raise HTTPException(403, f"Requires the '{perm}' permission on this project.")
    return user


def _release_read(r: LibraryRelease | None) -> ReleaseRead | None:
    if r is None:
        return None
    return ReleaseRead(
        id=r.id,
        library_id=r.library_id,
        name=r.name,
        notes=r.notes,
        item_count=len(r.item_versions or {}),
        released_by=r.released_by.display_name if r.released_by else None,
        released_at=r.released_at,
        previous_release_id=r.previous_release_id,
    )


def _record_read(db: Session, r: ApplicabilityRecord, params: dict, detail: bool = False):
    snap = lib.snapshot_at(db, r.source_item_id, r.source_version) or {}
    title, missing_t = lib.resolve_text(snap.get("title"), params)
    desc, missing_d = lib.resolve_text(snap.get("description_md"), params)
    copy = r.modified_item
    base = dict(
        id=r.id,
        library_id=r.library_id,
        release_name=r.release.name,
        source_item_id=r.source_item_id,
        source_version=r.source_version,
        key=snap.get("key"),
        title=snap.get("title") or "",
        description_md=snap.get("description_md"),
        resolved_title=title,
        resolved_description=desc,
        missing_parameters=sorted(set(missing_t + missing_d)),
        subsystem=r.subsystem,
        status=r.status,
        status_label=wf.state_label(SYSTEM_WORKFLOWS["applicability"], r.status),
        rationale=r.rationale,
        modified_item_id=r.modified_item_id,
        modified_item_key=copy.key if copy else None,
        modified_document_id=copy.document_id if copy else None,
        decided_by=r.decided_by.display_name if r.decided_by else None,
        decided_at=r.decided_at,
        approved_by=r.approved_by.display_name if r.approved_by else None,
        approved_at=r.approved_at,
        suspect=r.suspect,
        suspect_reason=r.suspect_reason,
        removed_from_library=r.removed_from_library,
        suggested_status=r.suggested_status,
        suggested_reason=r.suggested_reason,
    )
    if detail:
        return RecordDetail(**base, modification_diff=lib.modification_diff(db, r),
                            attributes=snap.get("attributes") or {})
    return RecordRead(**base)


def _subscription_read(db: Session, sub: ProjectLibrary) -> SubscriptionRead:
    latest = lib.latest_release(db, sub.library_id)
    update = bool(latest and latest.id != sub.release_id and latest.id != sub.declined_release_id
                  and latest.released_at > sub.release.released_at)
    return SubscriptionRead(
        library_id=sub.library_id,
        library_name=sub.library.name,
        release=_release_read(sub.release),
        latest_release=_release_read(latest),
        update_available=update,
        declined_release=_release_read(sub.declined_release),
        declined_justification=sub.declined_justification,
        target_document_id=sub.target_document_id,
    )


# ---- libraries & releases ---------------------------------------------------------


@router.get("/organizations/{org_id}/libraries", response_model=list[LibraryRead])
def list_libraries(org_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)):
    libs = db.query(Project).filter(Project.org_id == org_id, Project.kind == "library").order_by(Project.name).all()
    out = []
    for p in libs:
        if not can_read(db, p, user.id if user else None):
            continue
        out.append(
            LibraryRead(
                id=p.id,
                name=p.name,
                latest_release=_release_read(lib.latest_release(db, p.id)),
                requirement_count=len(lib._content_items(db, p.id)),
                mission_attribute_defs=p.mission_attribute_defs or [],
            )
        )
    return out


@router.get("/projects/{library_id}/releases", response_model=list[ReleaseRead])
def list_releases(library_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)):
    _project(db, library_id, user)
    rows = (
        db.query(LibraryRelease)
        .filter(LibraryRelease.library_id == library_id)
        .order_by(LibraryRelease.released_at.desc())
        .all()
    )
    return [_release_read(r) for r in rows]


@router.post("/projects/{library_id}/releases", response_model=ReleaseRead, status_code=201)
def create_release(
    library_id: uuid.UUID, payload: ReleaseCreate, user: User = Depends(require_user), db: Session = Depends(get_db)
):
    library = _project(db, library_id, user)
    if not (has_org_permission(db, user, "library.release") or has_permission(db, library_id, user, "project.admin")):
        raise HTTPException(403, "Requires the 'library.release' permission (or admin on the library).")
    try:
        release = lib.create_release(db, library, payload.name.strip(), payload.notes, user)
    except lib.LibraryError as e:
        raise HTTPException(422, str(e))
    db.commit()
    db.refresh(release)
    return _release_read(release)


@router.get("/library-releases/{release_id}/diff")
def release_diff(
    release_id: uuid.UUID,
    against: uuid.UUID | None = None,
    user: User | None = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Differences of `release_id` vs `against` (default: its previous release)."""
    release = db.get(LibraryRelease, release_id)
    if not release:
        raise HTTPException(404, "Release not found")
    _project(db, release.library_id, user)
    old_id = against or release.previous_release_id
    old = db.get(LibraryRelease, old_id) if old_id else None
    return lib.release_diff(db, old, release)


@router.put("/projects/{library_id}/mission-attribute-defs")
def set_mission_defs(
    library_id: uuid.UUID, payload: MissionDefsIn, user: User = Depends(require_user), db: Session = Depends(get_db)
):
    library = _project(db, library_id, user)
    _require(db, library_id, user, "project.admin")
    for d in payload.defs:
        if not d.get("key") or d.get("type", "boolean") not in ("boolean", "enum", "number", "text"):
            raise HTTPException(422, "Each attribute needs a key and a type (boolean, enum, number, text).")
    library.mission_attribute_defs = payload.defs
    db.commit()
    return {"defs": library.mission_attribute_defs}


@router.get("/items/{item_id}/library-rules")
def get_rules(item_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)):
    item = db.get(Item, item_id)
    if not item:
        raise HTTPException(404, "Item not found")
    _project(db, item.document.project_id, user)
    row = db.query(LibraryRule).filter_by(item_id=item_id).first()
    return {"rules": row.rules if row else []}


@router.put("/items/{item_id}/library-rules")
def set_rules(item_id: uuid.UUID, payload: RulesIn, user: User = Depends(require_user), db: Session = Depends(get_db)):
    item = db.get(Item, item_id)
    if not item:
        raise HTTPException(404, "Item not found")
    library = _project(db, item.document.project_id, user)
    if library.kind != "library":
        raise HTTPException(422, "Applicability rules belong to library requirements.")
    _require(db, library.id, user, "item.edit")
    for rule in payload.rules:
        if rule.get("op", "eq") not in lib._OPS or not rule.get("attribute"):
            raise HTTPException(422, f"Rule needs an attribute and one of the operators {sorted(lib._OPS)}.")
        if rule.get("suggest") not in ("applicable", "not_applicable"):
            raise HTTPException(422, "A rule can suggest 'applicable' or 'not_applicable'.")
    row = db.query(LibraryRule).filter_by(item_id=item_id).first()
    if row is None:
        row = LibraryRule(item_id=item_id, rules=payload.rules)
        db.add(row)
    else:
        row.rules = payload.rules
    db.commit()
    return {"rules": row.rules}


# ---- project side ---------------------------------------------------------------------


@router.get("/projects/{project_id}/libraries", response_model=list[SubscriptionRead])
def list_subscriptions(project_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)):
    _project(db, project_id, user)
    return [_subscription_read(db, s) for s in db.query(ProjectLibrary).filter_by(project_id=project_id).all()]


@router.post("/projects/{project_id}/libraries", response_model=SubscriptionRead, status_code=201)
def subscribe(project_id: uuid.UUID, payload: SubscribeIn, user: User = Depends(require_user), db: Session = Depends(get_db)):
    project = _project(db, project_id, user)
    _require(db, project_id, user, "project.admin")
    library = _project(db, payload.library_id, user)
    release = db.get(LibraryRelease, payload.release_id)
    if not release:
        raise HTTPException(404, "Release not found")
    try:
        sub = lib.subscribe(db, project, library, release, user, payload.target_document_id)
    except lib.LibraryError as e:
        raise HTTPException(422, str(e))
    db.commit()
    db.refresh(sub)
    return _subscription_read(db, sub)


def _sub(db: Session, project_id: uuid.UUID, library_id: uuid.UUID) -> ProjectLibrary:
    sub = db.query(ProjectLibrary).filter_by(project_id=project_id, library_id=library_id).first()
    if not sub:
        raise HTTPException(404, "This project doesn't use that library.")
    return sub


@router.patch("/projects/{project_id}/libraries/{library_id}", response_model=SubscriptionRead)
def patch_subscription(
    project_id: uuid.UUID, library_id: uuid.UUID, payload: SubscriptionPatch,
    user: User = Depends(require_user), db: Session = Depends(get_db),
):
    _project(db, project_id, user)
    _require(db, project_id, user, "project.admin")
    sub = _sub(db, project_id, library_id)
    sub.target_document_id = payload.target_document_id
    db.commit()
    return _subscription_read(db, sub)


@router.get("/projects/{project_id}/libraries/{library_id}/update-preview")
def update_preview(
    project_id: uuid.UUID, library_id: uuid.UUID, release_id: uuid.UUID,
    user: User | None = Depends(get_current_user), db: Session = Depends(get_db),
):
    _project(db, project_id, user)
    sub = _sub(db, project_id, library_id)
    new = db.get(LibraryRelease, release_id)
    if not new or new.library_id != library_id:
        raise HTTPException(404, "Release not found")
    return lib.release_diff(db, sub.release, new)


@router.post("/projects/{project_id}/libraries/{library_id}/adopt")
def adopt(
    project_id: uuid.UUID, library_id: uuid.UUID, payload: AdoptIn,
    user: User = Depends(require_user), db: Session = Depends(get_db),
):
    project = _project(db, project_id, user)
    _require(db, project_id, user, "project.admin")
    sub = _sub(db, project_id, library_id)
    release = db.get(LibraryRelease, payload.release_id)
    if not release:
        raise HTTPException(404, "Release not found")
    try:
        diff = lib.adopt(db, project, sub, release, user)
    except lib.LibraryError as e:
        raise HTTPException(422, str(e))
    db.commit()
    return {"subscription": _subscription_read(db, sub), "diff": diff}


@router.post("/projects/{project_id}/libraries/{library_id}/decline", response_model=SubscriptionRead)
def decline(
    project_id: uuid.UUID, library_id: uuid.UUID, payload: DeclineIn,
    user: User = Depends(require_user), db: Session = Depends(get_db),
):
    project = _project(db, project_id, user)
    _require(db, project_id, user, "project.admin")
    sub = _sub(db, project_id, library_id)
    release = db.get(LibraryRelease, payload.release_id)
    if not release:
        raise HTTPException(404, "Release not found")
    try:
        lib.decline(db, project, sub, release, payload.justification, user)
    except lib.LibraryError as e:
        raise HTTPException(422, str(e))
    db.commit()
    return _subscription_read(db, sub)


@router.put("/projects/{project_id}/mission-attributes")
def set_mission_attributes(
    project_id: uuid.UUID, payload: MissionAttributesIn, user: User = Depends(require_user), db: Session = Depends(get_db)
):
    project = _project(db, project_id, user)
    _require(db, project_id, user, "project.admin")
    project.mission_attributes = payload.values
    suggested = lib.apply_rules(db, project)
    audit.log(db, action="project.mission_attributes", subject_type="project", subject_id=project.id, actor=user,
              project_id=project.id, summary="Mission attributes updated", payload=payload.values)
    db.commit()
    return {"values": project.mission_attributes, "suggestions": suggested}


@router.get("/projects/{project_id}/applicability", response_model=list[RecordRead])
def list_records(
    project_id: uuid.UUID,
    library_id: uuid.UUID | None = None,
    status: str | None = None,
    subsystem: str | None = None,
    suspect: bool | None = None,
    user: User | None = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _project(db, project_id, user)
    q = db.query(ApplicabilityRecord).filter(ApplicabilityRecord.project_id == project_id)
    if library_id:
        q = q.filter(ApplicabilityRecord.library_id == library_id)
    if status:
        q = q.filter(ApplicabilityRecord.status.in_(status.split(",")))
    if subsystem:
        q = q.filter(ApplicabilityRecord.subsystem == subsystem)
    if suspect is not None:
        q = q.filter(ApplicabilityRecord.suspect.is_(suspect))
    params = lib.parameter_values(db, project_id)
    rows = [_record_read(db, r, params) for r in q.all()]
    return sorted(rows, key=lambda r: (r.subsystem, r.key or ""))


@router.get("/applicability/{record_id}", response_model=RecordDetail)
def get_record(record_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)):
    r = db.get(ApplicabilityRecord, record_id)
    if not r:
        raise HTTPException(404, "Record not found")
    _project(db, r.project_id, user)
    return _record_read(db, r, lib.parameter_values(db, r.project_id), detail=True)


@router.post("/applicability/{record_id}/clear-suspect", response_model=RecordRead)
def clear_record_suspect(record_id: uuid.UUID, user: User = Depends(require_user), db: Session = Depends(get_db)):
    r = db.get(ApplicabilityRecord, record_id)
    if not r:
        raise HTTPException(404, "Record not found")
    _require(db, r.project_id, user, "applicability.decide")
    audit.log(db, action="applicability.suspect_cleared", subject_type="applicability", subject_id=r.id, actor=user,
              project_id=r.project_id, summary="Decision re-confirmed after library change",
              payload={"reason": r.suspect_reason})
    r.suspect = False
    r.suspect_reason = None
    db.commit()
    return _record_read(db, r, lib.parameter_values(db, r.project_id))


@router.get("/projects/{project_id}/applicability/summary")
def applicability_summary(project_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)):
    _project(db, project_id, user)
    return lib.summary(db, project_id)


_STATUS_TEXT = {
    "to_be_assessed": "To be assessed",
    "applicable": "Applicable",
    "applicable_modified": "Applicable with modification",
    "not_applicable_pending": "Not applicable (pending approval)",
    "not_applicable": "Not applicable",
}


@router.get("/projects/{project_id}/applicability/export.xlsx")
def export_xlsx(project_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)):
    project = _project(db, project_id, user)
    records = list_records(project_id, None, None, None, None, user, db)
    summary = lib.summary(db, project_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Applicability"
    header = ["Subsystem", "ID", "Requirement", "Library release", "Version", "Applicability", "Justification",
              "Project requirement", "Decided by", "Decided at", "Approved by", "Flags"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
    for r in records:
        flags = []
        if r.suspect:
            flags.append(f"Suspect: {r.suspect_reason}")
        if r.missing_parameters:
            flags.append("Missing parameters: " + ", ".join(r.missing_parameters))
        ws.append([
            r.subsystem, r.key or "", r.resolved_title + ("\n" + r.resolved_description if r.resolved_description else ""),
            r.release_name, r.source_version, _STATUS_TEXT.get(r.status, r.status), r.rationale or "",
            r.modified_item_key or "", r.decided_by or "", r.decided_at.strftime("%Y-%m-%d") if r.decided_at else "",
            r.approved_by or "", "; ".join(flags),
        ])
    for i, w in enumerate([14, 12, 60, 14, 8, 22, 40, 14, 14, 11, 14, 30], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Summary")
    statuses = list(_STATUS_TEXT)
    ws2.append(["Subsystem"] + [_STATUS_TEXT[s] for s in statuses] + ["Total"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for subsystem, counts in sorted(summary["by_subsystem"].items()):
        ws2.append([subsystem] + [counts.get(s, 0) for s in statuses] + [sum(counts.values())])
    ws2.append(["Total"] + [summary["totals"].get(s, 0) for s in statuses] + [summary["total"]])
    ws2.append([])
    ws2.append(["Matrix complete", "Yes" if summary["complete"] else "No"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Applicability-{project.name}.xlsx".replace(" ", "_")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/projects/{project_id}/applicability/export.reqif")
def export_reqif(project_id: uuid.UUID, user: User | None = Depends(get_current_user), db: Session = Depends(get_db)):
    from app.services.reqif import applicability_reqif

    project = _project(db, project_id, user)
    records = list_records(project_id, None, None, None, None, user, db)
    xml = applicability_reqif(project, records)
    filename = f"Applicability-{project.name}.reqif".replace(" ", "_")
    return Response(
        xml,
        media_type="application/xml",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
