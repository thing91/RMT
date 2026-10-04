"""Glossary import/export — deliberately mirrors the requirements export
pipeline (§8): same XLSX-via-openpyxl / HTML-via-string approach, so the
glossary can be handed to the same people/process as a requirements
document rather than needing a separate tool.
"""

import html
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.models.glossary import GlossaryTerm

_HEADERS = ["Abbreviation", "Term", "Definition", "Category"]


def build_glossary_workbook(terms: list[GlossaryTerm]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Glossary"
    ws.append(_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for term in sorted(terms, key=lambda t: t.abbreviation):
        ws.append([term.abbreviation, term.term, term.definition, term.category or ""])
    for i in range(1, len(_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 30
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_glossary_html(org_name: str, terms: list[GlossaryTerm]) -> str:
    rows = "".join(
        f"<tr><td><code>[[{html.escape(t.abbreviation)}]]</code></td>"
        f"<td>{html.escape(t.term)}</td><td>{html.escape(t.definition)}</td>"
        f"<td>{html.escape(t.category or '')}</td></tr>"
        for t in sorted(terms, key=lambda t: t.abbreviation)
    )
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Glossary — {html.escape(org_name)}</title>
<style>
body {{ font-family: system-ui, sans-serif; max-width: 900px; margin: 40px auto; padding: 0 20px; color: #1e293b; }}
table {{ border-collapse: collapse; width: 100%; }}
th, td {{ text-align: left; padding: 8px 12px; border-bottom: 1px solid #e2e8f0; vertical-align: top; }}
th {{ background: #f8fafc; font-size: 0.8rem; text-transform: uppercase; color: #64748b; }}
code {{ font-family: ui-monospace, monospace; background: #f1f5f9; padding: 1px 6px; border-radius: 4px; }}
</style>
</head>
<body>
<h1>Glossary — {html.escape(org_name)}</h1>
<table>
<thead><tr><th>Syntax</th><th>Term</th><th>Definition</th><th>Category</th></tr></thead>
<tbody>{rows}</tbody>
</table>
</body>
</html>
"""


class GlossaryImportError(ValueError):
    pass


def parse_glossary_workbook(file_bytes: bytes) -> list[dict]:
    """Returns a list of {abbreviation, term, definition, category} dicts.
    Header row is required and matched case-insensitively; column order
    doesn't matter as long as the four names are present.
    """
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        raise GlossaryImportError(f"Could not read this file as XLSX: {e}") from e

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise GlossaryImportError("The file is empty.")

    normalized = [str(h).strip().lower() if h else "" for h in header]
    try:
        col = {
            "abbreviation": normalized.index("abbreviation"),
            "term": normalized.index("term"),
            "definition": normalized.index("definition"),
        }
    except ValueError:
        raise GlossaryImportError(
            "Header row must include 'Abbreviation', 'Term', and 'Definition' columns."
        )
    category_col = normalized.index("category") if "category" in normalized else None

    results = []
    for row in rows_iter:
        if row is None or all(v is None for v in row):
            continue
        abbreviation = row[col["abbreviation"]]
        if not abbreviation:
            continue
        results.append(
            {
                "abbreviation": str(abbreviation).strip(),
                "term": str(row[col["term"]] or "").strip(),
                "definition": str(row[col["definition"]] or "").strip(),
                "category": (
                    str(row[category_col]).strip()
                    if category_col is not None and row[category_col]
                    else None
                ),
            }
        )
    return results
