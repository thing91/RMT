"""§8.2 XLSX exporter — one row per DocumentNode, structural rows styled
distinctly. Rich-text fields are flattened to plain text — the same
limitation ReqView itself has, per §8.2. Links are listed as plain text
too (no in-sheet jump, unlike the HTML export's anchors) since a link may
point outside this document/sheet entirely.
"""

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.item import Item
from app.models.item_type import ItemType
from app.services.document_tree import DocumentNode, attribute_columns, build_document_tree


def _strip_markdown(md: str | None) -> str:
    if not md:
        return ""
    return re.sub(r"[#*_`>-]", "", md).strip()


def _format_attribute(value) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return str(value)


def _format_links(node: DocumentNode) -> str:
    return "; ".join(f"{link.label} {link.target_key or link.target_title}" for link in node.links)


def build_document_workbook(
    db: Session, document_name: str, items: list[Item], item_types: list[ItemType]
) -> io.BytesIO:
    nodes = build_document_tree(db, items, item_types)
    columns = attribute_columns(item_types)

    wb = Workbook()
    ws = wb.active
    ws.title = document_name[:31] or "Document"

    header = ["#", "Key", "Title", "Description"] + [label for _, label in columns] + ["Links"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for node in nodes:
        if node.kind == "structural":
            row = [node.number or "", "", node.title, ""] + [""] * len(columns) + [""]
        else:
            row = (
                ["", node.key or "", node.title, _strip_markdown(node.description_md)]
                + [_format_attribute(node.attributes.get(key)) for key, _ in columns]
                + [_format_links(node)]
            )
        ws.append(row)
        if node.kind == "structural":
            row_idx = ws.max_row
            for col_idx in range(1, len(header) + 1):
                ws.cell(row=row_idx, column=col_idx).font = Font(bold=True)

    for col_idx, _ in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 24

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
